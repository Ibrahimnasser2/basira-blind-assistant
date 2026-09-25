# -*- coding: utf-8 -*-
"""Simple Word doc with just the AI / ML model code (no formal report)."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).parent
ARABIC_RE = re.compile(r"[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]")
LATIN_RE = re.compile(r"[A-Za-z]")


SECTIONS = [
    {
        "title": "النموذج الأول: التعرف على العملة المصرية – TFLite / YOLOv8",
        "desc": (
            "نموذج محلي بصيغة TFLite مبني على YOLOv8 لاكتشاف فئات العملة "
            "المصرية (5، 10، 20، 50، 100، 200 جنيه)، مع تحميل النموذج، "
            "معالجة الصورة، تنفيذ الاستدلال، واختيار أعلى احتمال."
        ),
        "file": ROOT / "lib/features/currency_detection/currency_screen.dart",
    },
    {
        "title": "النموذج الثاني: خدمة Gemini متعدد الوسائط",
        "desc": (
            "خدمة تستدعي نموذج Gemini 1.5 Flash لوصف الصور واستخراج "
            "النصوص وترجمة التسميات إلى العربية."
        ),
        "file": ROOT / "lib/core/services/gemini_service.dart",
    },
    {
        "title": "النموذج الثالث: اكتشاف الأشياء – ML Kit Object Detection و Image Labeling",
        "desc": (
            "يستخدم نماذج Google ML Kit لاكتشاف الأشياء وتصنيفها، مع طبقة "
            "ترجمة محلية إلى العربية واحتياطي ترجمة آلية."
        ),
        "file": ROOT / "lib/features/object_detection/object_screen.dart",
    },
    {
        "title": "النموذج الرابع: قراءة النصوص OCR – ML Kit مع Gemini كبديل",
        "desc": (
            "يعتمد على Google ML Kit Text Recognition للنصوص اللاتينية، "
            "ويستخدم Gemini كبديل للنصوص العربية والمختلطة."
        ),
        "file": ROOT / "lib/features/text_reader/ocr_screen.dart",
    },
    {
        "title": "النموذج الخامس: اكتشاف الوجوه – ML Kit Face Detection",
        "desc": "يستخدم نموذج Face Detection من Google ML Kit لعدّ الوجوه في الإطار.",
        "file": ROOT / "lib/features/face_recognition/face_screen.dart",
    },
    {
        "title": "النموذج السادس: وصف المشهد عبر Gemini",
        "desc": (
            "شاشة التقاط دورية للإطارات وإرسالها لنموذج Gemini "
            "لوصف المشهد صوتياً."
        ),
        "file": ROOT / "lib/features/scene_description/scene_screen.dart",
    },
    {
        "title": "النموذج السابع: التعرف على اللون الغالب",
        "desc": (
            "خوارزمية بسيطة لاستخراج اللون الغالب من الإطار "
            "وتحويله لاسم لون عربي."
        ),
        "file": ROOT / "lib/features/color_detection/color_screen.dart",
    },
]


# ---------------------------------------------------------------------------
# Bidi helpers
# ---------------------------------------------------------------------------
def split_bidi(text: str):
    if not text:
        return []
    out, buf, cur = [], [], None
    for ch in text:
        if ARABIC_RE.match(ch):
            d = "rtl"
        elif LATIN_RE.match(ch):
            d = "ltr"
        else:
            d = None
        if d is None:
            buf.append(ch)
            continue
        if cur is None:
            cur = d
            buf.append(ch)
        elif d == cur:
            buf.append(ch)
        else:
            out.append(("".join(buf), cur))
            buf, cur = [ch], d
    if buf:
        out.append(("".join(buf), cur or "rtl"))
    return out


def set_run_fonts(run, latin: str, cs: str) -> None:
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:cs"), cs)


def mark_run_rtl(run) -> None:
    rPr = run._element.get_or_add_rPr()
    rtl = rPr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        rPr.append(rtl)
    rtl.set(qn("w:val"), "1")


def set_rtl_paragraph(p) -> None:
    pPr = p._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
    bidi.set(qn("w:val"), "1")


def shade_paragraph(p, fill_hex: str) -> None:
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    pPr.append(shd)


def add_arabic(doc: Document, text: str, *, bold: bool = False,
               size: int = 12, color: RGBColor | None = None,
               align=WD_ALIGN_PARAGRAPH.JUSTIFY) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    set_rtl_paragraph(p)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    set_run_fonts(run, "Calibri", "Arial")
    mark_run_rtl(run)


def add_centered_arabic(doc: Document, text: str, *, size: int,
                        bold: bool = False, italic: bool = False,
                        color: RGBColor | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_rtl_paragraph(p)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    set_run_fonts(run, "Calibri", "Arial")
    mark_run_rtl(run)


def add_heading_ar(doc: Document, text: str, *, size: int = 14,
                   color: RGBColor | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl_paragraph(p)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    set_run_fonts(run, "Calibri", "Arial")
    mark_run_rtl(run)


def add_code_block(doc: Document, code: str) -> None:
    for line in code.splitlines() or [""]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.2)
        shade_paragraph(p, "F4F4F4")
        run = p.add_run(line if line else " ")
        run.font.size = Pt(9)
        set_run_fonts(run, "Consolas", "Consolas")


def add_file_caption(doc: Document, relpath: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(f"// File: {relpath}")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x70, 0x70, 0x70)
    set_run_fonts(run, "Consolas", "Consolas")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main() -> None:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(2.0)

    add_centered_arabic(
        doc, "أكواد نماذج الذكاء الاصطناعي – تطبيق بصيرة",
        size=18, bold=True, color=RGBColor(0x1F, 0x49, 0x7D),
    )
    add_centered_arabic(
        doc, "نسخة مبسّطة تحتوي فقط على كود الموديلات الذكية",
        size=12, italic=True, color=RGBColor(0x55, 0x55, 0x55),
    )

    add_arabic(
        doc,
        "يحتوي هذا الملف على كود الموديلات والخدمات الذكية المستخدمة "
        "داخل التطبيق فقط: نموذج TFLite YOLOv8 للعملة، ونماذج Google "
        "ML Kit للوجوه والنصوص والأشياء، ونموذج Gemini متعدد الوسائط "
        "لوصف المشهد وقراءة النصوص. تم استبعاد بقية كود الواجهات "
        "وملفات الإعداد التي لا علاقة لها بالنماذج الذكية.",
        size=11,
    )

    for section in SECTIONS:
        doc.add_paragraph()
        add_heading_ar(doc, section["title"], size=14,
                       color=RGBColor(0x1F, 0x49, 0x7D))
        add_arabic(doc, section["desc"], size=11)

        rel = section["file"].relative_to(ROOT).as_posix()
        add_file_caption(doc, rel)
        code = section["file"].read_text(encoding="utf-8")
        add_code_block(doc, code)

    out = ROOT / "Basira_AI_Models_Code.docx"
    doc.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
